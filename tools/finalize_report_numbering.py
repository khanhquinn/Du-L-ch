from pathlib import Path
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


def norm(text):
    return " ".join((text or "").split())


def insert_after(paragraph, text):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.add_run(text)
    return new_para


def main():
    path = Path(sys.argv[1])
    doc = Document(path)

    replacements = {
        "3.3.3 Đặt tour": "3.3.5 Đặt tour",
        "3.3.4 Gửi email xác nhận đặt tour": "3.3.6 Gửi email xác nhận đặt tour",
        "3.3.5 Đăng nhập Admin": "3.3.7 Đăng nhập Admin",
        "3.3.6 Quản lý booking": "3.3.8 Quản lý booking",
        "3.3.7 Quản lý lịch tour": "3.3.9 Quản lý lịch tour",
        "3.3.8 Khoá ngày": "3.3.10 Khoá ngày",
        "3.3.9 Mở khoá ngày": "3.3.11 Mở khoá ngày",
        "3.3.10 Xem thống kê": "3.3.12 Xem thống kê",
    }

    added_table_list = False
    for paragraph in list(doc.paragraphs):
        text = norm(paragraph.text)
        if text in replacements:
            paragraph.text = replacements[text]
        if text == "Bảng 3. Mô tả Use Case Xem chi tiết tour" and not added_table_list:
            p1 = insert_after(paragraph, "Bảng 4a. Mô tả Use Case Đăng kí/Đăng nhập khách hàng")
            insert_after(p1, "Bảng 4b. Mô tả Use Case Xem tour đã đặt")
            added_table_list = True

    doc.save(path)


if __name__ == "__main__":
    main()
