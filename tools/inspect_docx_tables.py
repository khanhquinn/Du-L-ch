from pathlib import Path
import sys

from docx import Document


def cell_text(cell):
    return " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())


def main():
    if len(sys.argv) < 2:
        raise SystemExit("Usage: inspect_docx_tables.py <docx>")

    doc = Document(Path(sys.argv[1]))
    for idx, table in enumerate(doc.tables):
        rows = []
        for row in table.rows[:4]:
            rows.append(" | ".join(cell_text(cell)[:100] for cell in row.cells))
        print(f"TABLE {idx}: {len(table.rows)} rows x {len(table.columns)} cols")
        for row in rows:
            print("  " + row)
        print()


if __name__ == "__main__":
    main()
