from pathlib import Path
import sys

from docx import Document


def main():
    if len(sys.argv) < 2:
      raise SystemExit("Usage: extract_docx_outline.py <docx>")

    path = Path(sys.argv[1])
    doc = Document(path)

    print("HEADINGS")
    for i, para in enumerate(doc.paragraphs):
        text = " ".join(para.text.split())
        if not text:
            continue
        style = para.style.name if para.style else ""
        if style.startswith("Heading") or style in {"Title", "Subtitle"}:
            print(f"{i:04d} | {style} | {text}")

    print("\nKEYWORD HITS")
    keywords = [
        "đăng nhập", "đăng ký", "đăng kí", "admin", "booking", "đặt tour",
        "email", "khách hàng", "chức năng", "sơ đồ", "lược đồ", "ca sử dụng",
    ]
    for i, para in enumerate(doc.paragraphs):
        text = " ".join(para.text.split())
        lower = text.lower()
        if text and any(k in lower for k in keywords):
            print(f"{i:04d} | {para.style.name if para.style else ''} | {text[:220]}")


if __name__ == "__main__":
    main()
