from __future__ import annotations

import argparse
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


FIG_RE = re.compile(r"^Hình\s+(\d+):\s*(.+)$")
TBL_RE = re.compile(r"^Bảng\s+(\d+):\s*(.+)$")


def clear_paragraph(p):
    for r in list(p.runs):
        p._p.remove(r._r)


def set_text(p, text, *, bold=None, italic=None):
    clear_paragraph(p)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def text_of(p):
    return " ".join((p.text or "").split())


def has_image(p):
    return bool(p._p.xpath(".//w:drawing")) or bool(p._p.xpath(".//w:pict"))


def remove_element(el):
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def body_blocks(doc):
    pi = ti = 0
    out = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            out.append(("p", doc.paragraphs[pi], child))
            pi += 1
        elif tag == "tbl":
            out.append(("tbl", doc.tables[ti], child))
            ti += 1
        else:
            out.append((tag, None, child))
    return out


def find_para(doc, exact):
    for p in doc.paragraphs:
        if text_of(p) == exact:
            return p
    return None


def insert_paragraph_after(p, text="", style=None):
    new_p = OxmlElement("w:p")
    p._p.addnext(new_p)
    para = p._parent.paragraphs[-1]
    # python-docx cannot wrap by element directly; rebuild Document references by
    # using the paragraph inserted before the following save/load boundary.
    # The caller should reload if it needs object-level access. This helper is
    # intentionally not used for object-sensitive edits.
    return para


def add_paragraph_after_element(parent, after_el, text="", style=None):
    new_el = OxmlElement("w:p")
    after_el.addnext(new_el)
    p = parent.paragraphs[-1]
    # Returned wrapper may not point at new_el in python-docx internals, so do
    # direct OOXML text insertion.
    if style:
        ppr = new_el.get_or_add_pPr()
        pstyle = OxmlElement("w:pStyle")
        pstyle.set(qn("w:val"), style)
        ppr.append(pstyle)
    if text:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        r.append(t)
        new_el.append(r)
    return new_el


def element_text(el):
    return "".join(el.xpath(".//w:t/text()"))


def make_paragraph_el(text="", style_id=None):
    p = OxmlElement("w:p")
    if style_id:
        ppr = OxmlElement("w:pPr")
        pstyle = OxmlElement("w:pStyle")
        pstyle.set(qn("w:val"), style_id)
        ppr.append(pstyle)
        p.append(ppr)
    if text:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        r.append(t)
        p.append(r)
    return p


def add_page_break(p):
    run = p.add_run()
    run.add_break()


def set_style_by_text(doc):
    replacements = {
        "LỜI CẢM ƠN": ("LỜI CẢM ƠN", "Heading 1"),
        "CHƯƠNG1:GIỚI THIỆU": ("CHƯƠNG 1: GIỚI THIỆU", "Heading 1"),
        "CHƯƠNG 2: CƠ SỞ LÝ THUYẾT VÀ CÔNG NGHỆ SỬ DỤNG": (
            "CHƯƠNG 2: CƠ SỞ LÝ THUYẾT VÀ CÔNG NGHỆ SỬ DỤNG",
            "Heading 1",
        ),
        "Lý do chọn đề tài": ("1.1 Lý do chọn đề tài", "Heading 2"),
        "Mục tiêu đề tài": ("1.2 Mục tiêu đề tài", "Heading 2"),
        "Phạm vi đề tài": ("1.3 Phạm vi đề tài", "Heading 2"),
        "Phương pháp thực hiện": ("1.4 Phương pháp thực hiện", "Heading 2"),
        "Ý nghĩa khoa học và thực tiễn": ("1.5 Ý nghĩa khoa học và thực tiễn", "Heading 2"),
        "Cấu trúc báo cáo": ("1.6 Cấu trúc báo cáo", "Heading 2"),
        "2.1 Tổng quan về Website": ("2.1 Tổng quan về Website", "Heading 2"),
    }
    for p in doc.paragraphs:
        t = text_of(p)
        if t in replacements:
            new_text, style = replacements[t]
            set_text(p, new_text, bold=True)
            p.style = style

    # Two paragraphs in section 2.1 were imported as List Paragraph even
    # though they are prose.
    for p in doc.paragraphs:
        t = text_of(p)
        if t.startswith("Website du lịch là hệ thống trực tuyến") or t.startswith(
            "Bên cạnh vai trò quảng bá hình ảnh du lịch"
        ):
            p.style = "Normal"


def add_urgency_paragraph(doc):
    anchor = None
    for p in doc.paragraphs:
        if text_of(p).startswith("Thông qua đề tài này, em có cơ hội vận dụng"):
            anchor = p
            break
    if anchor is None:
        return

    new_el = deepcopy(anchor._p)
    for node in list(new_el):
        new_el.remove(node)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = (
        "Bên cạnh yêu cầu về giao diện, đề tài còn có tính cấp thiết ở việc lựa chọn "
        "công nghệ triển khai phù hợp với xu hướng phát triển ứng dụng web hiện nay. "
        "Node.js giúp xây dựng các API xử lý nghiệp vụ đặt tour theo mô hình bất đồng "
        "bộ, phù hợp với các thao tác truy xuất dữ liệu và phản hồi nhanh cho người dùng. "
        "Vercel hỗ trợ triển khai frontend và serverless API trực tiếp từ mã nguồn, "
        "giúp rút ngắn thời gian đưa sản phẩm lên môi trường thật. Railway cung cấp "
        "dịch vụ PostgreSQL trực tuyến, thuận lợi cho việc lưu trữ dữ liệu tour, lịch "
        "khởi hành và booking. Việc kết hợp các công nghệ này giúp hệ thống có khả năng "
        "vận hành online, dễ bảo trì và có thể mở rộng trong tương lai."
    )
    r.append(t)
    new_el.append(r)
    anchor._p.addnext(new_el)


def replace_database_section(doc):
    blocks = body_blocks(doc)
    start = end = None
    for i, (kind, obj, el) in enumerate(blocks):
        if kind == "p" and text_of(obj) == "3.4.1 Sơ đồ ERD":
            start = i
        if kind == "p" and text_of(obj) == "3.4.3 Sơ đồ luồng xử lý đặt tour":
            end = i
            break
    if start is None or end is None:
        return

    heading = blocks[start][1]
    set_text(heading, "3.4.1 Lược đồ quan hệ giữa các bảng", bold=True)
    heading.style = "Heading 3"

    for _kind, _obj, el in blocks[start + 1 : end]:
        remove_element(el)

    insert_after = heading._p
    new_texts = [
        (
            "Cơ sở dữ liệu PostgreSQL của hệ thống được thiết kế theo mô hình quan hệ, "
            "trong đó các bảng chính liên kết với nhau thông qua khóa chính và khóa ngoại. "
            "Cách trình bày lược đồ quan hệ giúp thể hiện rõ sự phụ thuộc dữ liệu giữa các "
            "bảng mà không lặp lại phần mô tả chi tiết từng cột."
        ),
        "Lược đồ quan hệ chính của hệ thống:",
        "destinations(id) 1 - N tours(destination_id): một điểm đến có thể gắn với nhiều tour.",
        "tours(id) 1 - N tour_schedules(tour_id): một tour có nhiều lịch khởi hành.",
        "tours(id) 1 - N bookings(tour_id): một tour có thể phát sinh nhiều đơn đặt tour.",
        "customers(id) 1 - N bookings(customer_id): một khách hàng có thể tạo nhiều đơn đặt tour.",
        "blocked_dates(id): lưu các ngày bị khóa chung hoặc khóa theo từng tour.",
        "admins(id): lưu tài khoản quản trị viên, phục vụ chức năng đăng nhập và quản lý hệ thống.",
    ]
    for text in reversed(new_texts):
        new_p = make_paragraph_el(text)
        heading._p.addnext(new_p)


def fix_database_following_heading(doc):
    for p in doc.paragraphs:
        if text_of(p) == "3.4.3 Sơ đồ luồng xử lý đặt tour":
            set_text(p, "3.4.2 Sơ đồ luồng xử lý đặt tour", bold=True)
            p.style = "Heading 3"


def delete_old_front_matter(doc):
    abbrev_rows = []
    if doc.tables:
        try:
            for row in doc.tables[0].rows:
                abbrev_rows.append([cell.text for cell in row.cells])
        except Exception:
            pass

    blocks = body_blocks(doc)
    start = end = None
    for i, (kind, obj, _el) in enumerate(blocks):
        if kind == "p" and text_of(obj) == "DANH SÁCH HÌNH ẢNH":
            start = i
        if kind == "p" and text_of(obj) == "CHƯƠNG 1: GIỚI THIỆU":
            end = i
            break
    if start is not None and end is not None:
        for _kind, _obj, el in blocks[start:end]:
            remove_element(el)
    return abbrev_rows


def renumber_captions(doc):
    for p in doc.paragraphs:
        t = text_of(p)
        m = FIG_RE.match(t)
        if m:
            n = int(m.group(1))
            caption = m.group(2).replace("Postgres databasez", "Postgres database")
            if n > 8:
                set_text(p, f"Hình {n - 1}: {caption}")
            elif n == 8 and "ERD" in caption:
                # This caption should have been removed with the ERD section.
                pass
            else:
                set_text(p, f"Hình {n}: {caption}")
        m = TBL_RE.match(t)
        if m:
            n = int(m.group(1))
            caption = m.group(2)
            if n >= 18:
                set_text(p, f"Bảng {n - 7}: {caption}")


def move_captions(doc):
    body = doc.element.body

    def tag(el):
        return el.tag.split("}")[-1]

    def is_blank_para(el):
        return tag(el) == "p" and not element_text(el).strip() and not (
            bool(el.xpath(".//w:drawing")) or bool(el.xpath(".//w:pict"))
        )

    def is_image_para(el):
        return tag(el) == "p" and (
            bool(el.xpath(".//w:drawing")) or bool(el.xpath(".//w:pict"))
        )

    changed = True
    while changed:
        changed = False
        children = list(body.iterchildren())
        for i, el in enumerate(children):
            if tag(el) != "p" or not FIG_RE.match(" ".join(element_text(el).split())):
                continue
            j = i + 1
            while j < len(children) and is_blank_para(children[j]):
                j += 1
            if j < len(children) and is_image_para(children[j]):
                remove_element(el)
                children[j].addnext(el)
                changed = True
                break

    changed = True
    while changed:
        changed = False
        children = list(body.iterchildren())
        for i, el in enumerate(children):
            if tag(el) != "tbl":
                continue
            j = i + 1
            while j < len(children) and is_blank_para(children[j]):
                j += 1
            if j < len(children) and tag(children[j]) == "p":
                txt = " ".join(element_text(children[j]).split())
                if TBL_RE.match(txt):
                    cap_el = children[j]
                    remove_element(cap_el)
                    el.addprevious(cap_el)
                    changed = True
                    break
    return doc


def collect_captions(doc):
    figs, tbls = [], []
    for p in doc.paragraphs:
        t = text_of(p)
        if FIG_RE.match(t):
            figs.append(t)
        elif TBL_RE.match(t):
            tbls.append(t)
    return figs, tbls


def fix_special_caption_order(doc):
    body = doc.element.body

    def tag(el):
        return el.tag.split("}")[-1]

    def txt(el):
        return " ".join(element_text(el).split())

    def is_img(el):
        return tag(el) == "p" and (
            bool(el.xpath(".//w:drawing")) or bool(el.xpath(".//w:pict"))
        )

    def find_para(prefix):
        for el in body.iterchildren():
            if tag(el) == "p" and txt(el).startswith(prefix):
                return el
        return None

    def find_next_image(after_el):
        seen = False
        for el in body.iterchildren():
            if el is after_el:
                seen = True
                continue
            if seen and is_img(el):
                return el
        return None

    def find_next_para(after_el, prefix):
        seen = False
        for el in body.iterchildren():
            if el is after_el:
                seen = True
                continue
            if seen and tag(el) == "p" and txt(el).startswith(prefix):
                return el
        return None

    # Textual directory tree: place the figure label below the tree block.
    cap = find_para("Hình 9: Cấu trúc thư mục dự án")
    if cap is not None:
        target = None
        for el in body.iterchildren():
            if tag(el) == "p" and txt(el).startswith("Du Lich/"):
                target = el
                break
        if target is not None:
            remove_element(cap)
            target.addnext(cap)

    # A few imported screenshots had a description between caption and image.
    for prefix in ["Hình 14: Trang chủ", "Hình 17: Trang đặt tour"]:
        cap = find_para(prefix)
        if cap is not None:
            img = find_next_image(cap)
            if img is not None:
                remove_element(cap)
                img.addnext(cap)

    # Responsive section has two consecutive images and two consecutive labels.
    cap34 = find_para("Hình 34: Mobile responsive")
    cap35 = find_para("Hình 35: Mobile responsive")
    if cap34 is not None and cap35 is not None:
        images_before = []
        for el in body.iterchildren():
            if el is cap34:
                break
            if is_img(el):
                images_before.append(el)
        if len(images_before) >= 2:
            img34, img35 = images_before[-2], images_before[-1]
            remove_element(cap35)
            img35.addnext(cap35)
            remove_element(cap34)
            img34.addnext(cap34)

    # Remove an empty imported Heading 3 paragraph that was previously entering
    # the document outline.
    for p in list(doc.paragraphs):
        if p.style.name.startswith("Heading") and not text_of(p):
            remove_element(p._p)


def insert_front_matter(doc, abbrev_rows):
    chapter = find_para(doc, "CHƯƠNG 1: GIỚI THIỆU")
    if chapter is None:
        return
    figs, tbls = collect_captions(doc)
    chapter_titles = []
    for p in doc.paragraphs:
        t = text_of(p)
        if not t:
            continue
        if p.style.name in {"Heading 1", "Heading 2", "Heading 3"}:
            if t not in {"MỤC LỤC"}:
                chapter_titles.append((p.style.name, t))

    toc_entries = []
    for style, title in chapter_titles:
        if title == "LỜI CẢM ƠN":
            toc_entries.append((style, title))
            break
    toc_entries.extend(
        [
            ("Heading 1", "DANH MỤC HÌNH ẢNH"),
            ("Heading 1", "DANH MỤC BẢNG"),
            ("Heading 1", "DANH MỤC TỪ VIẾT TẮT"),
        ]
    )
    toc_entries.extend([(style, title) for style, title in chapter_titles if title != "LỜI CẢM ƠN"])

    insert_at = chapter._p
    items = []
    items.append(("Heading1", "MỤC LỤC"))
    for style, t in toc_entries:
        level = {"Heading 1": 1, "Heading 2": 2, "Heading 3": 3}.get(style, 1)
        items.append(("TOC", ("    " * (level - 1)) + t))
    items.append(("PAGEBREAK", ""))
    items.append(("Heading1", "DANH MỤC HÌNH ẢNH"))
    for t in figs:
        items.append(("Normal", t))
    items.append(("PAGEBREAK", ""))
    items.append(("Heading1", "DANH MỤC BẢNG"))
    for t in tbls:
        items.append(("Normal", t))
    items.append(("PAGEBREAK", ""))
    items.append(("Heading1", "DANH MỤC TỪ VIẾT TẮT"))

    for style, content in items:
        if style == "PAGEBREAK":
            p = make_paragraph_el()
            r = OxmlElement("w:r")
            br = OxmlElement("w:br")
            br.set(qn("w:type"), "page")
            r.append(br)
            p.append(r)
        elif style == "TOC":
            p = make_paragraph_el(content)
        else:
            p = make_paragraph_el(content, style)
        insert_at.addprevious(p)

    # Insert a clean abbreviation table before Chapter 1, after the new title.
    tmp = Path.cwd() / "_tmp_front.docx"
    doc.save(str(tmp))
    doc = Document(str(tmp))
    title = find_para(doc, "DANH MỤC TỪ VIẾT TẮT")
    if title and abbrev_rows:
        table = doc.add_table(rows=0, cols=len(abbrev_rows[0]))
        for row_vals in abbrev_rows:
            row = table.add_row()
            for cell, val in zip(row.cells, row_vals):
                cell.text = val
        title._p.addnext(table._tbl)
        pb = make_paragraph_el()
        r = OxmlElement("w:r")
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        r.append(br)
        pb.append(r)
        table._tbl.addnext(pb)
    tmp.unlink(missing_ok=True)
    return doc


def set_cell_text_style(cell, bold=False, align=None):
    for p in cell.paragraphs:
        if align is not None:
            p.alignment = align
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        for run in p.runs:
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(12)
            run.bold = bold or run.bold


def set_cell_margins(cell, top=90, right=120, bottom=90, left=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in {"top": top, "right": right, "bottom": bottom, "left": left}.items():
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_width(table, widths):
    tblPr = table._tbl.tblPr
    tblLayout = tblPr.find(qn("w:tblLayout"))
    if tblLayout is None:
        tblLayout = OxmlElement("w:tblLayout")
        tblPr.append(tblLayout)
    tblLayout.set(qn("w:type"), "fixed")
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    total = sum(widths)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcW = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"), str(widths[min(idx, len(widths) - 1)]))
            tcW.set(qn("w:type"), "dxa")


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    hdr = trPr.find(qn("w:tblHeader"))
    if hdr is None:
        hdr = OxmlElement("w:tblHeader")
        trPr.append(hdr)
    hdr.set(qn("w:val"), "true")


def format_tables(doc):
    section = doc.sections[0]
    usable = int(section.page_width.twips - section.left_margin.twips - section.right_margin.twips)
    for table in doc.tables:
        table.style = "Table Grid"
        cols = len(table.columns)
        if cols == 2:
            widths = [int(usable * 0.32), usable - int(usable * 0.32)]
        elif cols == 3:
            widths = [int(usable * 0.16), int(usable * 0.28), usable - int(usable * 0.44)]
        elif cols == 4:
            widths = [int(usable * 0.12), int(usable * 0.40), int(usable * 0.18), usable - int(usable * 0.70)]
        else:
            widths = [int(usable / max(cols, 1))] * cols
        set_table_width(table, widths)
        if table.rows:
            repeat_header(table.rows[0])
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                align = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 or c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                set_cell_text_style(cell, bold=(r_idx == 0), align=align)
                if r_idx == 0:
                    shade_cell(cell, "D9EAF7")


def format_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)

    for style_name, size, bold in [
        ("Normal", 13, False),
        ("Normal (Web)", 13, False),
        ("Heading 1", 14, True),
        ("Heading 2", 13, True),
        ("Heading 3", 13, True),
    ]:
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = bold

    for p in doc.paragraphs:
        t = text_of(p)
        if not t and not has_image(p):
            continue
        for run in p.runs:
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            if run.font.size is None:
                run.font.size = Pt(13)
        pf = p.paragraph_format
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        if p.style.name == "Heading 1":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = None
            pf.space_before = Pt(12)
            pf.space_after = Pt(12)
        elif p.style.name in {"Heading 2", "Heading 3"}:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.first_line_indent = None
            pf.space_before = Pt(8)
            pf.space_after = Pt(6)
        elif FIG_RE.match(t):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = None
            for run in p.runs:
                run.italic = True
                run.font.size = Pt(12)
        elif TBL_RE.match(t):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = None
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(12)
        elif has_image(p):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = None
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.first_line_indent = Cm(1.0)


def format_front_matter_lists(doc):
    active = None
    front_titles = {"DANH MỤC HÌNH ẢNH", "DANH MỤC BẢNG", "MỤC LỤC"}
    for p in doc.paragraphs:
        t = text_of(p)
        if t in front_titles:
            active = t
            continue
        if t == "DANH MỤC TỪ VIẾT TẮT":
            active = None
            continue
        if active and p.style.name == "Heading 1":
            active = None
        if active and t:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = None
            p.paragraph_format.left_indent = Cm(0.0)
            p.paragraph_format.line_spacing = 1.2
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                run.bold = False
                run.italic = False
                run.font.size = Pt(13)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    doc = Document(args.input)
    set_style_by_text(doc)
    add_urgency_paragraph(doc)
    replace_database_section(doc)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)

    doc = Document(args.output)
    fix_database_following_heading(doc)
    renumber_captions(doc)
    abbrev_rows = delete_old_front_matter(doc)
    doc.save(args.output)

    doc = Document(args.output)
    doc = move_captions(doc)
    fix_special_caption_order(doc)
    doc.save(args.output)

    doc = Document(args.output)
    doc = insert_front_matter(doc, abbrev_rows)
    format_document(doc)
    format_front_matter_lists(doc)
    format_tables(doc)
    doc.save(args.output)
    print(f"Saved {args.output}")


if __name__ == "__main__":
    main()
